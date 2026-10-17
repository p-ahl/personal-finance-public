#!/usr/bin/env python3
"""Charts of the household finances over time: cash vs goal, income vs spending, assets and
liabilities, spending by category. Every point is a query over the ledger (v_ledger); each month
is reconciled to the cent against v_balance_sheet and v_income_statement_monthly before anything
is written (the script exits non-zero if a check fails).

  python3 scripts/charts.py --as-of 2026-10-10 [--db ../finance-data/finance.db]
      [--out ../finance-data/exports/charts-2026-10-10] [--no-png]

Points are month ends from the first booked month, plus the as-of date (month to date). Lines
dated after --as-of (future accruals) are left out. Outputs: charts.html (inline SVG, no external
scripts; follows prefers-color-scheme, tap a column to read its values), PNGs in light and dark
(matplotlib), charts.xlsx (the series), validation.txt. No account numbers are printed.
Re-running with the same inputs rewrites the same files.
"""
import argparse
import html
import json
import os
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reportcfg as RC  # noqa: E402  (names, series labels, notes: $FINANCE_DATA/config/report.local.yaml)
import statements as S  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.normpath(os.path.join(HERE, "..", "..", "finance-data"))
CFG = RC.load()
TRIPS = (CFG["trips_accounts"] or [""])[0]

# cash series: (key, label, account prefixes). Order = stacking order.
CASH_SERIES = [(x["key"], x["label"], tuple(x["prefixes"])) for x in CFG["cash_series"]]
ASSET_GROUPS = [("cash", "Cash"), ("inv", "Investments"), ("crypto", "Crypto"), ("re", "Real estate"),
                ("veh", "Vehicles"), ("oth_a", "Other assets")]
LOAN = CFG["loan_group"]
LIAB_GROUPS = [("mort", "Mortgage")] + ([("land", LOAN["label"])] if LOAN["account_prefixes"] else []) + \
              [("cards", "Credit cards"), ("pay", "Payables"), ("tax", "Accrued taxes")]
COLORS = {**{x["key"]: x.get("color", "#4e79a7") for x in CFG["cash_series"]},
          "trips": "#b07aa1", "cash": "#4e79a7", "inv": "#59a14f", "crypto": "#f28e2b", "re": "#9c755f",
          "veh": "#76b7b2", "oth_a": "#bab0ac", "mort": "#e15759", "land": "#ff9d9a", "cards": "#d37295",
          "pay": "#c49c94", "tax": "#edc948", "salary": "#4e79a7", "business": "#59a14f", "partner": "#76b7b2",
          "realized": "#9c6ade", "other_inc": "#bab0ac", "living": "#e15759", "taxes": "#edc948",
          "goal": "#e15759", "net": "#2f8f5b", "nw": "#2f8f5b", "total": "#4e79a7"}
INCOME_SERIES = [(x["key"], x["label"], x["prefix"]) for x in CFG["income_series"]]
INCOME = [(k, lab) for k, lab, _ in INCOME_SERIES] + [("realized", "Investment gains sold"), ("other_inc", "Other income")]


def money(c):
    s = f"${abs(c) / 100:,.0f}"
    return f"-{s}" if c < 0 else s


def money_c(c):
    s = f"${abs(c) / 100:,.2f}"
    return f"-{s}" if c < 0 else s


def month_end(y, mo):
    return (date(y + mo // 12, mo % 12 + 1, 1) - timedelta(days=1))


def asset_group(a):
    if a.startswith("Assets:Bank:"):
        return "cash"
    if a.startswith("Assets:Investments:"):
        return "inv"
    if a.startswith("Assets:Crypto:"):
        return "crypto"
    if a.startswith("Assets:RealEstate:"):
        return "re"
    if a.startswith("Assets:Vehicles:"):
        return "veh"
    return "oth_a"


def liab_group(a):
    if a.startswith("Liabilities:Mortgage"):
        return "mort"
    if any(a.startswith(p) for p in LOAN["account_prefixes"]):
        return "land"
    if a.startswith("Liabilities:CreditCard"):
        return "cards"
    if a.startswith("Liabilities:Accrued:Taxes"):
        return "tax"
    return "pay"


def income_key(a):
    for k, _, prefix in INCOME_SERIES:
        if a.startswith(prefix):
            return k
    if a.startswith("Income:Investments:Unrealized"):
        return "unrealized"
    if a.startswith("Income:Investments:Realized"):
        return "realized"
    return "other_inc"


# ---------------------------------------------------------------- data

def collect(db, as_of, trips=TRIPS):
    c = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    q = lambda sql, *a: c.execute(sql, a).fetchall()  # noqa: E731
    ao = as_of.isoformat()
    first = q("select min(date) from v_ledger")[0][0]
    y, mo = int(first[:4]), int(first[5:7])
    points = []  # (label, date)
    while date(y, mo, 1) <= as_of:
        me = month_end(y, mo)
        if me < as_of:
            points.append((f"{me:%b %y}", me))
        else:
            points.append((f"{as_of:%b} {as_of.day} (MTD)", as_of))
        y, mo = (y + 1, 1) if mo == 12 else (y, mo + 1)
    rows = q("select date, account, amount_minor from v_ledger where account_type in ('asset','liability') and date <= ?", ao)
    cash_accts = {r[0] for r in q("select distinct account from v_ledger where cash_flow_class = 'cash'")}

    def bal_on(dt):
        b = defaultdict(int)
        s = dt.isoformat()
        for d_, a, v in rows:
            if d_ <= s:
                b[a] += v
        return b

    D = {"as_of": ao, "labels": [p[0] for p in points], "dates": [p[1].isoformat() for p in points]}
    cash = {k: [] for k, _, _ in CASH_SERIES}
    cash.update(trips=[], goal_cash=[], total=[])
    ag = {k: [] for k, _ in ASSET_GROUPS}
    lg = {k: [] for k, _ in LIAB_GROUPS}
    nw, assets_t, liab_t = [], [], []
    unclassified_cash = set()
    for _, dt in points:
        b = bal_on(dt)
        tot = 0
        for k, _, pre in CASH_SERIES:
            v = sum(x for a, x in b.items() if a in cash_accts and a != trips and a.startswith(pre))
            cash[k].append(v)
        cash["trips"].append(b.get(trips, 0))
        tot = sum(x for a, x in b.items() if a in cash_accts)
        cash["total"].append(tot)
        cash["goal_cash"].append(tot - b.get(trips, 0))
        for a in b:
            if a in cash_accts and a != trips and not any(a.startswith(p) for _, _, p2 in CASH_SERIES for p in p2):
                unclassified_cash.add(a)
        for k, _ in ASSET_GROUPS:
            ag[k].append(sum(x for a, x in b.items() if a.startswith("Assets:") and asset_group(a) == k))
        for k, _ in LIAB_GROUPS:
            lg[k].append(-sum(x for a, x in b.items() if a.startswith("Liabilities:") and liab_group(a) == k))
        assets_t.append(sum(x for a, x in b.items() if a.startswith("Assets:")))
        liab_t.append(-sum(x for a, x in b.items() if a.startswith("Liabilities:")))
        nw.append(assets_t[-1] - liab_t[-1])
    assert not unclassified_cash, f"cash accounts outside the cash series: {unclassified_cash}"
    D.update(cash=cash, assets=ag, liabs=lg, nw=nw, assets_total=assets_t, liab_total=liab_t)

    # weekly cash (Fridays) for the HTML
    wk_labels, wk_goal, wk_trips = [], [], []
    d0 = date.fromisoformat(first)
    fri = d0 + timedelta(days=(4 - d0.weekday()) % 7)
    while fri <= as_of:
        b = bal_on(fri)
        t = b.get(trips, 0)
        wk_labels.append(fri.isoformat())
        wk_goal.append(sum(x for a, x in b.items() if a in cash_accts) - t)
        wk_trips.append(t)
        fri += timedelta(days=7)
    D["weekly"] = {"labels": wk_labels, "goal_cash": wk_goal, "trips": wk_trips}

    # income vs spending by month (dates <= as_of)
    pl = q("""select month, account, account_type, level2, natural_minor from v_ledger
              where statement = 'income_statement' and date <= ?""", ao)
    months = [p[1].isoformat()[:7] for p in points]
    inc = {k: [0] * len(months) for k, _ in INCOME}
    inc["unrealized"] = [0] * len(months)
    living, taxes = [0] * len(months), [0] * len(months)
    cats = defaultdict(lambda: [0] * len(months))
    for mth, a, typ, lvl2, v in pl:
        i = months.index(mth)
        if typ == "income":
            inc[income_key(a)][i] += v
        elif a.startswith("Expenses:Taxes"):
            taxes[i] += v
        else:
            living[i] += v
            cats[lvl2][i] += v
    income_t = [sum(inc[k][i] for k, _ in INCOME) for i in range(len(months))]
    net = [income_t[i] - living[i] - taxes[i] for i in range(len(months))]
    D.update(months=months, income=inc, income_total=income_t, living=living, taxes=taxes, net_savings=net,
             cats={k: v for k, v in cats.items()})

    # annotations (only from ledger text)
    notes = defaultdict(list)
    for pat, lab in [(x["pattern"], x["label"]) for x in CFG["income_notes"]]:
        for mth, s in q("""select month, sum(natural_minor) from v_ledger where statement = 'income_statement'
                           and account_type = 'expense' and description like ? and date <= ? group by month""", pat, ao):
            if s and abs(s) >= 100000:
                notes[mth].append(f"{lab} {money(s)}")
    for mth, s in q("""select month, sum(natural_minor) from v_ledger where account like 'Expenses:Travel:Vacation%'
                       and date <= ? group by month having sum(natural_minor) >= 300000""", ao):
        notes[mth].append(f"Vacation {money(s)}")
    D["inc_notes"] = {k: v for k, v in notes.items()}

    # valuation-method events (for the balance-sheet chart)
    ev = []
    for dt, desc, s in q("""select date, min(description), sum(amount_minor) from v_ledger where source = 'valuation'
                            and account_type = 'asset' and date <= ? group by date, entry_id order by date""", ao):
        kind = "Crypto" if "elf-custody" in desc else "Investments"
        what = ("self-custody crypto re-stated at cost" if "at-cost" in desc else
                "self-custody crypto switched to market value" if kind == "Crypto" else "investments marked to market")
        amt = "net $0 (moved between wallets)" if s == 0 else ("+" if s > 0 else "") + money_c(s)
        ev.append({"date": dt, "label": f"{dt}: {what}, {amt}", "amount": s})
    D["valuation_events"] = ev
    D["opening_date"] = q("select min(date) from v_ledger where source = 'opening' and date <= ?", ao)[0][0]

    D["checks"] = validate(q, D, points, ao)
    return D


def validate(q, D, points, ao):
    """Every chart point vs the reporting views, to the cent. Raises on any mismatch."""
    out, bad = [], []
    # balance sheet: v_balance_sheet is all-time, so remove lines dated after each point from it
    bs_all = dict(q("""select section, sum(amount_minor) from v_balance_sheet where section in ('asset','liability')
                       group by section"""))
    for i, (_, dt) in enumerate(points):
        after = dict(q("""select account_type, sum(natural_minor) from v_ledger where account_type in ('asset','liability')
                          and date > ? group by 1""", dt.isoformat()))
        exp_a = bs_all.get("asset", 0) - after.get("asset", 0)
        exp_l = bs_all.get("liability", 0) - after.get("liability", 0)
        got_a = sum(D["assets"][k][i] for k, _ in ASSET_GROUPS)
        got_l = sum(D["liabs"][k][i] for k, _ in LIAB_GROUPS)
        cash_v = sum(D["cash"][k][i] for k, _, _ in CASH_SERIES) + D["cash"]["trips"][i]
        ok = got_a == exp_a == D["assets_total"][i] and got_l == exp_l and cash_v == D["cash"]["total"][i] \
            and D["nw"][i] == exp_a - exp_l
        (out if ok else bad).append(f"{dt} balance sheet: assets {money_c(got_a)} vs {money_c(exp_a)}, "
                                    f"liabilities {money_c(got_l)} vs {money_c(exp_l)} {'OK' if ok else 'MISMATCH'}")
    # income statement: v_income_statement_monthly minus lines in that month dated after as-of
    for i, mth in enumerate(D["months"]):
        view = dict(q("select section, sum(amount_minor) from v_income_statement_monthly where month = ? group by 1", mth))
        after = dict(q("""select account_type, sum(natural_minor) from v_ledger where statement = 'income_statement'
                          and month = ? and date > ? group by 1""", mth, ao))
        exp_i = view.get("income", 0) - after.get("income", 0)
        exp_e = view.get("expense", 0) - after.get("expense", 0)
        got_i = D["income_total"][i] + D["income"]["unrealized"][i]
        got_e = D["living"][i] + D["taxes"][i]
        ok = got_i == exp_i and got_e == exp_e and sum(v[i] for v in D["cats"].values()) == D["living"][i]
        (out if ok else bad).append(f"{mth} income statement: income {money_c(got_i)} vs {money_c(exp_i)}, "
                                    f"expenses {money_c(got_e)} vs {money_c(exp_e)} {'OK' if ok else 'MISMATCH'}")
    if bad:
        raise SystemExit("charts: validation failed\n" + "\n".join(bad))
    return out


# ---------------------------------------------------------------- inline SVG

def _nice_step(span):
    raw = span / 5 or 1
    p = 10 ** (len(str(int(raw))) - 1)
    for k in (1, 2, 2.5, 5, 10):
        if raw <= k * p:
            return k * p
    return 10 * p


def svg_chart(cid, labels, pos=(), neg=(), lines=(), hlines=(), marks=(), kind="bar", height=250, title="",
              tick_every=1):
    """pos/neg: [(name, color, values_cents)] stacked up / down. lines: [(name, color, values, dashed)].
    hlines: [(value_cents, label, color)]. marks: [(index, short label)]. Colors are literals; text and
    grid use the report's CSS variables so the chart follows light/dark."""
    W, H, L, R, T, B = 440, height, 52, 8, 14, 34
    n = len(labels)
    up = [sum(s[2][i] for s in pos) for i in range(n)]
    dn = [-sum(s[2][i] for s in neg) for i in range(n)]
    vals = up + dn + [v for l in lines for v in l[2]] + [h[0] for h in hlines] + [0]
    lo, hi = min(vals) / 100, max(vals) / 100
    st = _nice_step(hi - lo)
    lo, hi = (lo // st) * st, -((-hi) // st) * st
    if hi == lo:
        hi = lo + st
    pw, ph = W - L - R, H - T - B
    X = lambda i: L + pw * (i + 0.5) / n  # noqa: E731
    Y = lambda c: T + ph * (hi - c / 100) / (hi - lo)  # noqa: E731
    bw = pw / n * 0.7
    e = html.escape
    p = [f'<svg class="chart" id="{cid}" viewBox="0 0 {W} {H}" width="100%" role="img" aria-label="{e(title)}">']
    v = lo
    while v <= hi + 1e-9:
        y = Y(v * 100)
        lab = f"${v / 1000:,.0f}k" if abs(hi - lo) >= 5000 else f"${v:,.0f}"
        p.append(f'<line x1="{L}" x2="{W - R}" y1="{y:.1f}" y2="{y:.1f}" stroke="var(--line)" '
                 f'stroke-width="{1.4 if v == 0 else 0.8}"/>'
                 f'<text x="{L - 4}" y="{y + 4:.1f}" text-anchor="end" font-size="10" fill="var(--muted)">{lab}</text>')
        v += st
    for i, lab in enumerate(labels):
        if i % tick_every == 0 or i == n - 1:
            p.append(f'<text x="{X(i):.1f}" y="{H - B + 13}" text-anchor="middle" font-size="9" fill="var(--muted)">'
                     f'{e(lab.split(" (")[0] if n > 8 else lab)}</text>')

    def stack(series, sign):
        base = [0] * n
        for name, col, vs in series:
            if kind == "area":
                top = [base[i] + sign * vs[i] for i in range(n)]
                pts = [f"{X(i):.1f},{Y(top[i]):.1f}" for i in range(n)] + \
                      [f"{X(i):.1f},{Y(base[i]):.1f}" for i in reversed(range(n))]
                p.append(f'<polygon points="{" ".join(pts)}" fill="{col}" fill-opacity=".78" stroke="{col}" stroke-width=".6"/>')
                base = top
            else:
                for i in range(n):
                    if vs[i] == 0:
                        continue
                    y0, y1 = Y(base[i]), Y(base[i] + sign * vs[i])
                    p.append(f'<rect x="{X(i) - bw / 2:.1f}" y="{min(y0, y1):.1f}" width="{bw:.1f}" '
                             f'height="{abs(y1 - y0):.1f}" fill="{col}"/>')
                    base[i] += sign * vs[i]
    stack(pos, 1)
    stack(neg, -1)
    for val, lab, col in hlines:
        p.append(f'<line x1="{L}" x2="{W - R}" y1="{Y(val):.1f}" y2="{Y(val):.1f}" stroke="{col}" stroke-width="1.6" '
                 f'stroke-dasharray="6 4"/><text x="{W - R}" y="{Y(val) - 4:.1f}" text-anchor="end" font-size="10" '
                 f'fill="{col}" font-weight="600">{e(lab)}</text>')
    for name, col, vs, dashed in lines:
        pts = " ".join(f"{X(i):.1f},{Y(vs[i]):.1f}" for i in range(n))
        p.append(f'<polyline points="{pts}" fill="none" stroke="{col}" stroke-width="2.2"'
                 f'{" stroke-dasharray=" + chr(34) + "4 3" + chr(34) if dashed else ""}/>')
        if n <= 20:
            p += [f'<circle cx="{X(i):.1f}" cy="{Y(vs[i]):.1f}" r="2.4" fill="{col}"/>' for i in range(n)]
    for k, (i, lab) in enumerate(marks):
        p.append(f'<line x1="{X(i):.1f}" x2="{X(i):.1f}" y1="{T}" y2="{H - B}" stroke="var(--ink)" stroke-width="1" '
                 f'stroke-dasharray="2 3"/><text x="{X(i) - 3:.1f}" y="{T + 10 + 11 * k}" text-anchor="end" '
                 f'font-size="10" font-weight="700" fill="var(--ink)">{e(lab)}</text>')
    for i in range(n):  # tap targets
        p.append(f'<rect class="hit" data-i="{i}" x="{L + pw * i / n:.1f}" y="{T}" width="{pw / n:.1f}" height="{ph}" '
                 f'fill="transparent"/>')
    p.append("</svg>")
    leg = "".join(f'<span><i style="background:{s[1]}"></i>{e(s[0])}</span>' for s in list(pos) + list(neg)) + \
        "".join(f'<span><i class="ln" style="background:{l[1]}"></i>{e(l[0])}</span>' for l in lines) + \
        "".join(f'<span><i class="ln" style="background:{h[2]}"></i>{e(h[1])}</span>' for h in hlines)
    data = {"labels": labels, "series": [[s[0], s[2]] for s in list(pos) + list(neg) + list(lines)]}
    return (f'<div class="chartbox">{"".join(p)}<div class="legend">{leg}</div>'
            f'<div class="readout" id="{cid}-r">Tap a column to see its numbers.</div>'
            f'<script type="application/json" id="{cid}-d">{json.dumps(data)}</script></div>')


CHART_CSS = """
.chartbox { margin: 8px 0 4px; }
svg.chart { display: block; touch-action: manipulation; }
svg.chart .hit { cursor: pointer; }
svg.chart .hit.on { fill: var(--ink); fill-opacity: .07; }
.legend { display: flex; flex-wrap: wrap; gap: 4px 12px; font-size: 12px; color: var(--muted); margin: 4px 0; }
.legend i { display: inline-block; width: 10px; height: 10px; border-radius: 2px; margin-right: 4px; vertical-align: -1px; }
.legend i.ln { height: 3px; vertical-align: 2px; }
.readout { font-size: 13px; color: var(--muted); border-left: 3px solid var(--line); padding: 2px 8px; min-height: 20px; }
.readout b { color: var(--ink); }
"""

CHART_JS = """<script>
document.querySelectorAll('svg.chart').forEach(function (svg) {
  var d = JSON.parse(document.getElementById(svg.id + '-d').textContent), r = document.getElementById(svg.id + '-r');
  function f(c) { var s = '$' + (Math.abs(c) / 100).toLocaleString('en-US', {minimumFractionDigits: 2, maximumFractionDigits: 2});
    return c < 0 ? '-' + s : s; }
  svg.querySelectorAll('.hit').forEach(function (h) {
    h.addEventListener('click', function () {
      svg.querySelectorAll('.hit.on').forEach(function (x) { x.classList.remove('on'); });
      h.classList.add('on'); var i = +h.dataset.i;
      r.innerHTML = '<b>' + d.labels[i] + '</b>: ' + d.series.filter(function (s) { return s[1][i] !== 0; })
        .map(function (s) { return s[0] + ' ' + f(s[1][i]); }).join(' · ');
    });
  });
});
</script>"""


def cash_svg(D, cid="c-cash", goal=CFG["goal"], compact=False):
    s = [(lab, COLORS[k], D["cash"][k]) for k, lab, _ in CASH_SERIES if any(D["cash"][k])]
    return svg_chart(cid, D["labels"], pos=s, lines=[("Trips fund (not counted)", COLORS["trips"], D["cash"]["trips"], True)],
                     hlines=[(goal * 100, f"${goal / 1000:.0f}k goal", COLORS["goal"])],
                     height=210 if compact else 250, title="Cash counted toward goal, by bank, month ends",
                     tick_every=2 if compact else 1)


def income_svg(D, cid="c-inc", compact=False):
    pos = [(lab, COLORS[k], D["income"][k]) for k, lab in INCOME if any(D["income"][k])]
    neg = [("Living spending", COLORS["living"], D["living"]), ("Taxes", COLORS["taxes"], D["taxes"])]
    marks = [(D["months"].index(m), "★" * (j + 1)) for j, m in enumerate(sorted(D["inc_notes"]))] if not compact else []
    return svg_chart(cid, D["labels"], pos=pos, neg=neg, lines=[("Net savings", COLORS["net"], D["net_savings"], False)],
                     marks=marks, height=210 if compact else 270, title="Income vs spending by month",
                     tick_every=2 if compact else 1)


def bs_svg(D, cid="c-bs"):
    pos = [(lab, COLORS[k], D["assets"][k]) for k, lab in ASSET_GROUPS if any(D["assets"][k])]
    neg = [(lab, COLORS[k], D["liabs"][k]) for k, lab in LIAB_GROUPS if any(D["liabs"][k])]
    marks = []
    if D["valuation_events"]:
        marks.append((len(D["labels"]) - 1 if D["valuation_events"][-1]["date"][:7] == D["as_of"][:7]
                      else D["months"].index(D["valuation_events"][-1]["date"][:7]), "Valuation change"))
    return svg_chart(cid, D["labels"], pos=pos, neg=neg, lines=[("Net worth", COLORS["nw"], D["nw"], False)],
                     marks=marks, kind="area", height=300, title="Assets and liabilities")


def last6_cats(D):
    n = len(D["months"])
    idx = list(range(max(0, n - 6), n))
    tot = {k: sum(v[i] for i in idx) for k, v in D["cats"].items()}
    return idx, sorted(((k, t) for k, t in tot.items() if t), key=lambda x: -x[1])


def cats_svg(D):
    idx, ranked = last6_cats(D)
    top = ranked[:8]
    W, rowh = 440, 22
    mx = max(t for _, t in top)
    e = html.escape
    p = [f'<svg class="bars" viewBox="0 0 {W} {rowh * len(top) + 4}" width="100%" role="img" aria-label="Spending by category">']
    for j, (k, t) in enumerate(top):
        y = j * rowh
        w = (W - 200) * t / mx
        p.append(f'<text x="0" y="{y + 15}" font-size="12" fill="var(--ink)">{e(k)}</text>'
                 f'<rect x="110" y="{y + 4}" width="{w:.1f}" height="14" rx="2" fill="{COLORS["living"]}" fill-opacity=".8"/>'
                 f'<text x="{114 + w:.1f}" y="{y + 15}" font-size="11" fill="var(--muted)">{money(t)}</text>')
    p.append("</svg>")
    return "".join(p), idx, ranked


# ---------------------------------------------------------------- HTML

def build_html(D, goal, theme=""):
    T = S.Terms({**{k: list(v) for k, v in S.GLOSSARY.items()}, **GLOSS}, interactive=not theme)
    e = html.escape
    n = len(D["labels"])
    c, i_ = D["cash"], n - 1
    out = [f'<!doctype html><html lang="en" class="{theme}"><head><meta charset="utf-8">'
           '<meta name="viewport" content="width=device-width, initial-scale=1">'
           f'<title>Charts {D["as_of"]}</title><style>{S.CSS}{CHART_CSS}</style></head><body><main>'
           f'<h1>Finances over time</h1><p class="sub">{e(CFG["household"])} · {e(D["labels"][0])} – {e(D["labels"][-1])} · '
           f'every point is from the books; month-end balances, last point is {D["as_of"]} (month to date)</p>']
    # 1 cash
    out.append(f'<section><h2>1. {T("[[cash|Cash]]")} over time</h2>' + cash_svg(D, goal=goal))
    out.append('<ul class="notes">'
               f'<li>Counted toward the ${goal:,} goal now: <b>{money_c(c["goal_cash"][i_])}</b> '
               f'({c["goal_cash"][i_] / goal:.0f}%).</li>'
               f'<li>{e(CFG["trips_chart_label"])} {money_c(c["trips"][i_])} is drawn as a dashed line and '
               'is not counted.</li><li>Books start with opening balances on '
               f'{D["opening_date"]}, so the first point is the first month end.</li></ul>')
    wk = D["weekly"]
    if len(wk["labels"]) > 4:
        out.append('<p class="sub" style="margin-top:12px">Week by week (Fridays)</p>' +
                   svg_chart("c-wk", [x[5:] for x in wk["labels"]], lines=[("Counted cash", COLORS["total"], wk["goal_cash"], False),
                             ("Trips fund", COLORS["trips"], wk["trips"], True)],
                             hlines=[(goal * 100, f"${goal / 1000:.0f}k goal", COLORS["goal"])], height=190,
                             title="Weekly cash", tick_every=max(1, len(wk["labels"]) // 6)))
    out.append("</section>")
    # 2 income vs spending
    out.append(f'<section><h2>2. Income vs spending</h2>' + income_svg(D))
    out.append('<ul class="notes"><li>Bars above zero are income; below zero are living spending and taxes '
               '(taxes kept apart). The line is net savings = income − spending − taxes.</li>'
               f'<li>{T("[[unrealized|Unrealized]]")} market gains and losses are left out of income '
               f'(total in the period: {money_c(sum(D["income"]["unrealized"]))}).</li>'
               '<li>Months are on the accrual basis, the same as the monthly income statement: a bill counts in the '
               'month the service happened.</li>')
    for j, mth in enumerate(sorted(D["inc_notes"])):
        out.append(f'<li>{"★" * (j + 1)} {mth}: {e("; ".join(D["inc_notes"][mth]))}</li>')
    out.append("</ul></section>")
    # 3 balance sheet
    out.append(f'<section><h2>3. Assets and liabilities</h2>' + bs_svg(D))
    out.append(f'<ul class="notes"><li>Net worth now {money_c(D["nw"][-1])}: assets {money_c(D["assets_total"][-1])}, '
               f'owed {money_c(D["liab_total"][-1])}.</li>'
               f'<li>{e(CFG["real_estate_note"])}</li>')
    for ev in D["valuation_events"]:
        out.append(f'<li><b>Valuation change</b> {e(ev["label"])}. This is a change in how the asset is measured, '
                   'not money in or out.</li>')
    out.append("</ul></section>")
    # 4 categories
    svg, idx, ranked = cats_svg(D)
    out.append(f'<section><h2>4. Spending by category</h2><p class="sub">{e(D["labels"][idx[0]])} – '
               f'{e(D["labels"][idx[-1]])}, living spending (taxes excluded), top 8</p>{svg}</section>')
    out.append(f'<section><h2>Checks</h2><p class="sub">{len(D["checks"])} month checks against the balance sheet '
               'and income statement views: all match to the cent.</p></section>')
    out.append('<details class="gloss"><summary>Glossary</summary>' + T.glossary_html() + "</details>")
    out.append(f"</main>{'<script>' + S.TIP_JS + '</script>' if not theme else ''}{CHART_JS}</body></html>")
    return "".join(out)


GLOSS = {
    "cash": ("Cash", "Money in bank accounts you can spend right away (checking, savings, Venmo). Investments and "
             "crypto are not cash until sold."),
    "unrealized": ("Unrealized gain or loss", "The change in market value of something still owned. Nothing was "
                   "sold, so it can still move back."),
}


# ---------------------------------------------------------------- PNG (matplotlib)

def pngs(D, out, goal):
    import matplotlib
    matplotlib.use("Agg")
    matplotlib.rcParams["svg.hashsalt"] = "fixed"
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    fmt = FuncFormatter(lambda v, _: ("-" if v < 0 else "") + f"${abs(v) / 1000:,.0f}k")
    x = list(range(len(D["labels"])))
    files = []
    for theme in ("light", "dark"):
        fg, bg, grid = (("#1b1f24", "#ffffff", "#e3e6ea") if theme == "light" else ("#ecebe8", "#111214", "#2a2e34"))
        plt.rcParams.update({"text.color": fg, "axes.labelcolor": fg, "xtick.color": fg, "ytick.color": fg,
                             "axes.edgecolor": grid, "figure.facecolor": bg, "axes.facecolor": bg, "font.size": 9})

        def fig(title):
            f, ax = plt.subplots(figsize=(6, 4.2), dpi=160)
            ax.set_title(title, loc="left", fontweight="bold")
            ax.yaxis.set_major_formatter(fmt)
            ax.grid(axis="y", color=grid, lw=.7)
            ax.set_axisbelow(True)
            ax.set_xticks(x)
            ax.set_xticklabels([l.replace(" (MTD)", "\nMTD") for l in D["labels"]], rotation=45, ha="right", fontsize=7)
            for s in ("top", "right"):
                ax.spines[s].set_visible(False)
            return f, ax

        def save(f, ax, name):
            ax.legend(fontsize=7, frameon=False, ncol=2, loc="upper left")
            f.tight_layout()
            p = os.path.join(out, f"{name}-{theme}.png")
            f.savefig(p, metadata={"Software": None})
            plt.close(f)
            files.append(p)

        # 1 cash
        f, ax = fig(f"Cash toward ${goal // 1000}k goal (trips fund excluded)")
        base = [0] * len(x)
        for k, lab, _ in CASH_SERIES:
            v = [c / 100 for c in D["cash"][k]]
            if any(v):
                ax.bar(x, v, bottom=base, color=COLORS[k], label=lab, width=.7)
                base = [b + w for b, w in zip(base, v)]
        ax.plot(x, [c / 100 for c in D["cash"]["trips"]], "--o", ms=3, color=COLORS["trips"], label="Trips fund (not counted)")
        ax.axhline(goal, color=COLORS["goal"], ls="--", lw=1.4, label=f"${goal // 1000}k goal")
        save(f, ax, "1-cash")
        # 2 income vs spending
        f, ax = fig("Income vs spending by month (taxes separate)")
        base = [0] * len(x)
        for k, lab in INCOME:
            v = [c / 100 for c in D["income"][k]]
            if any(v):
                ax.bar(x, v, bottom=base, color=COLORS[k], label=lab, width=.7)
                base = [b + w for b, w in zip(base, v)]
        lv = [-c / 100 for c in D["living"]]
        ax.bar(x, lv, color=COLORS["living"], label="Living spending", width=.7)
        ax.bar(x, [-c / 100 for c in D["taxes"]], bottom=lv, color=COLORS["taxes"], label="Taxes", width=.7)
        ax.plot(x, [c / 100 for c in D["net_savings"]], "-o", ms=3, color=COLORS["net"], label="Net savings")
        ax.axhline(0, color=fg, lw=.8)
        lo_, hi_ = ax.get_ylim()
        ax.set_ylim(lo_ - (hi_ - lo_) * .14, hi_)
        for j, mth in enumerate(sorted(D["inc_notes"])):
            i = D["months"].index(mth)
            ax.annotate("\n".join(D["inc_notes"][mth]), (i, -(D["living"][i] + D["taxes"][i]) / 100), fontsize=6,
                        xytext=(0, -8), textcoords="offset points", ha="center" if i < len(x) - 3 else "right", va="top", color=fg)
        save(f, ax, "2-income-spending")
        # 3 balance sheet
        f, ax = fig("Assets, liabilities and net worth")
        ax.stackplot(x, *[[c / 100 for c in D["assets"][k]] for k, _ in ASSET_GROUPS],
                     labels=[l for _, l in ASSET_GROUPS], colors=[COLORS[k] for k, _ in ASSET_GROUPS], alpha=.8)
        ax.stackplot(x, *[[-c / 100 for c in D["liabs"][k]] for k, _ in LIAB_GROUPS],
                     labels=[l for _, l in LIAB_GROUPS], colors=[COLORS[k] for k, _ in LIAB_GROUPS], alpha=.8)
        ax.plot(x, [c / 100 for c in D["nw"]], "-o", ms=3, color=fg, label="Net worth")
        if D["valuation_events"]:
            ax.axvline(x[-1] if D["valuation_events"][-1]["date"][:7] == D["as_of"][:7] else
                       D["months"].index(D["valuation_events"][-1]["date"][:7]), color=fg, ls=":", lw=1)
            ax.text(x[-1] - .2, ax.get_ylim()[1] * .97, "valuation\nmethod\nchanged", ha="right", va="top", fontsize=6.5)
        save(f, ax, "3-assets-liabilities")
        # 4 categories
        idx, ranked = last6_cats(D)
        top = ranked[:8][::-1]
        f, ax = plt.subplots(figsize=(6, 3.6), dpi=160)
        ax.barh([k for k, _ in top], [t / 100 for _, t in top], color=COLORS["living"])
        ax.xaxis.set_major_formatter(fmt)
        ax.set_title(f"Spending by category, {D['labels'][idx[0]]} – {D['labels'][idx[-1]]}", loc="left", fontweight="bold")
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        f.tight_layout()
        p = os.path.join(out, f"4-categories-{theme}.png")
        f.savefig(p, metadata={"Software": None})
        plt.close(f)
        files.append(p)
    return files


# ---------------------------------------------------------------- xlsx

def build_xlsx(path, D):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, header, rows):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row[1:]:
                if isinstance(c.value, float):
                    c.number_format = '#,##0.00;[Red]-#,##0.00'
        ws.column_dimensions["A"].width = 16
        ws.freeze_panes = "B2"

    d = lambda c: round(c / 100, 2)  # noqa: E731
    n = range(len(D["labels"]))
    sheet("Cash", ["Date"] + [l for _, l, _ in CASH_SERIES] + ["Counted toward goal", "Trips fund", "Total cash"],
          [[D["dates"][i]] + [d(D["cash"][k][i]) for k, _, _ in CASH_SERIES] +
           [d(D["cash"]["goal_cash"][i]), d(D["cash"]["trips"][i]), d(D["cash"]["total"][i])] for i in n])
    sheet("Cash weekly", ["Friday", "Counted toward goal", "Trips fund"],
          [[w, d(g), d(t)] for w, g, t in zip(D["weekly"]["labels"], D["weekly"]["goal_cash"], D["weekly"]["trips"])])
    sheet("Income-Spending", ["Month"] + [l for _, l in INCOME] + ["Income total", "Living spending", "Taxes",
                                                                   "Net savings", "Unrealized (excluded)"],
          [[D["months"][i]] + [d(D["income"][k][i]) for k, _ in INCOME] +
           [d(D["income_total"][i]), d(D["living"][i]), d(D["taxes"][i]), d(D["net_savings"][i]),
            d(D["income"]["unrealized"][i])] for i in n])
    sheet("Assets-Liabilities", ["Date"] + [l for _, l in ASSET_GROUPS] + ["Assets total"] + [l for _, l in LIAB_GROUPS] +
          ["Liabilities total", "Net worth"],
          [[D["dates"][i]] + [d(D["assets"][k][i]) for k, _ in ASSET_GROUPS] + [d(D["assets_total"][i])] +
           [d(D["liabs"][k][i]) for k, _ in LIAB_GROUPS] + [d(D["liab_total"][i]), d(D["nw"][i])] for i in n])
    cats = sorted(D["cats"])
    sheet("Categories", ["Month"] + cats, [[D["months"][i]] + [d(D["cats"][k][i]) for k in cats] for i in n])
    sheet("Valuation events", ["Date", "What", "Asset change"],
          [[v["date"], v["label"], d(v["amount"])] for v in D["valuation_events"]])
    sheet("Checks", ["Check"], [[x] for x in D["checks"]])
    wb.save(path)


# ---------------------------------------------------------------- weekly report hook

def weekly_section(db, as_of, goal=CFG["goal"], trips=TRIPS):
    """Compact cash and income-vs-spending charts for weekly.py (HTML fragment incl. CSS and JS)."""
    D = collect(db, as_of, trips)
    return (f"<style>{CHART_CSS}</style>"
            '<p class="sub" style="margin-top:12px">Cash toward goal, month ends (trips fund dashed, not counted)</p>'
            + cash_svg(D, "w-cash", goal, compact=True) +
            '<p class="sub" style="margin-top:12px">Income vs spending by month (taxes separate; line = net savings)</p>'
            + income_svg(D, "w-inc", compact=True) + CHART_JS)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--as-of", required=True)
    ap.add_argument("--db", default=os.environ.get("FINANCE_DB", os.path.join(DATA, "finance.db")))
    ap.add_argument("--out", default="")
    ap.add_argument("--goal", type=int, default=CFG["goal"])
    ap.add_argument("--trips", default=TRIPS)
    ap.add_argument("--no-png", action="store_true")
    a = ap.parse_args()
    as_of = date.fromisoformat(a.as_of)
    out = a.out or os.path.join(DATA, "exports", f"charts-{as_of}")
    os.makedirs(out, exist_ok=True)
    D = collect(a.db, as_of, a.trips)
    html_s = build_html(D, a.goal)
    assert not re.search(r"(Assets|Liabilities):\S*-\d{3,}", html_s), "account number in output"
    open(os.path.join(out, "charts.html"), "w").write(html_s)
    xp = os.path.join(out, "charts.xlsx")
    build_xlsx(xp, D)
    S.reproducible_xlsx(xp, str(as_of))
    open(os.path.join(out, "validation.txt"), "w").write("\n".join(D["checks"]) + "\n")
    files = [] if a.no_png else pngs(D, out, a.goal)
    print(f"charts: {len(D['checks'])} checks OK; {len(D['labels'])} points; out {out}; {len(files)} PNGs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
