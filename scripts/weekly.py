#!/usr/bin/env python3
"""Weekly financial report: net worth bridge, the week's Profit and Loss, cash vs goal,
upcoming bills, open items, and trend charts (scripts/charts.py). Every figure is a query over the ledger (v_ledger and the accrual
views); estimates are labelled.

  python3 scripts/weekly.py --as-of 2026-10-10 --week-start 2026-10-03 \
      [--db ../finance-data/finance.db] [--out ../finance-data/exports/weekly-2026-10-10]

Week = week-start .. as-of minus 1 day (Sat..Fri). "Now" = books through the as-of date (the
nightly refresh books that morning's prices and pending items on the as-of date). Baseline =
balance sheet the day before week-start (or --baseline). With a prior report (--prior-report, default:
the newest earlier exports/weekly-* folder) the bridge starts from that report's *reported* net worth and
balances, so restatements (e.g. market revaluations that the ledger re-dates to the latest refresh) are
shown as their own lines instead of fake weekly swings. Trips money (accounts in --trips, the
trips-fund receivable) is shown but excluded from the cash goal.

Outputs: weekly.html (follows prefers-color-scheme, tap-to-open terms), weekly.pdf (light),
weekly-dark.pdf, weekly.xlsx. Reuses scripts/statements.py styling. Needs openpyxl and Chrome.
No account numbers are printed (trailing digit masks are stripped from account names).
"""
import argparse
import html
import os
import re
import sqlite3
import statistics
import subprocess
import sys
import tempfile
from collections import defaultdict
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reportcfg as RC  # noqa: E402  (names, goal and narrative notes: $FINANCE_DATA/config/report.local.yaml)
import statements as S  # noqa: E402  (CSS, TIP_JS, Terms, GLOSSARY, find_chrome, reproducible_*)

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.normpath(os.path.join(HERE, "..", "..", "finance-data"))
CFG = RC.load()

GLOSS = {k: list(v) for k, v in S.GLOSSARY.items()}
GLOSS.update({
    "bridge": ("Net worth bridge", "A walk from last week's net worth to this week's, line by line. The lines add up "
               "exactly to the change, so nothing is left unexplained."),
    "realized": ("Realized gain", "Profit locked in by selling an investment for more than it cost. Counts as income."),
    "unrealized": ("Unrealized gain or loss", "The change in market value of something still owned. Nothing was sold, "
                   "so it can still move back."),
    "cash": ("Cash and cash equivalents", "Money in bank accounts you can spend right away (checking, savings, Venmo). "
             "Investments and crypto are not cash until sold."),
    "trailing": ("Trailing 4-week average", "Average weekly spending over the 4 weeks just before this one. A quick "
                 "way to see whether this week was normal."),
    "pending": ("Pending", "Activity a bank or app shows but that has not posted yet. It is held in a temporary "
                "account until it posts, so it is not counted twice."),
    "uncategorized": ("Uncategorized", "A transaction the rules could not assign to a spending or income category. "
                      "It is still in the totals, just not in the right bucket."),
    "unmatched": ("Unmatched transfer", "Money that looks like it moved between our own accounts, but the other side "
                  "was not found. Usually a Venmo or check that needs a note."),
    "restatement": ("Restatement", "A change to an earlier period's figures after they were reported, for example "
                    "a transaction that posted late but is dated last week. Shown separately so it is not mistaken "
                    "for this week's activity."),
    "liquidity": ("Spendable cash", "Cash minus the credit card balance already owed; what is left if the card were "
                  "paid today."),
})

NICE_YEAR = None
TRIPS_DEFAULT = ",".join(CFG["trips_accounts"])
TRIPS_RECEIVABLE = CFG["trips_receivable_account"]


# ---------------------------------------------------------------- helpers

def d(s):
    return date.fromisoformat(s)


def nice(iso):
    x = d(iso) if isinstance(iso, str) else iso
    yr = f", {x.year}" if NICE_YEAR and x.year != NICE_YEAR else ""
    return f"{x:%a} {x:%b} {x.day}{yr}"


def m(cents, sign=False):
    """cents -> dollars with thousands separators; negatives in parentheses."""
    s = f"${abs(cents) / 100:,.2f}"
    if cents < 0:
        return f"({s})"
    return ("+" + s) if sign and cents > 0 else s


def strip_mask(name):
    return re.sub(r"-\d{3,}\b", "", name)


def load_labels(path):
    lab = {}
    if path and os.path.exists(path):
        for line in open(path):
            mt = re.match(r"\s+([A-Z][\w:.-]+):\s*(.+?)\s*$", line)
            if mt:
                lab[mt.group(1)] = mt.group(2)
    return lab


class Book:
    def __init__(self, db, labels, trips):
        self.c = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
        self.labels, self.trips = labels, trips

    def q(self, sql, *a):
        return self.c.execute(sql, a).fetchall()

    def label(self, acct):
        if acct in self.trips:
            base = self.labels.get(acct) or self._inst(acct) + " savings"
            return base + CFG["trips_fund_suffix"]
        mt = re.match(r"Liabilities:Accrued:Taxes:(" + "|".join(map(re.escape, CFG["tax_jurisdictions"])) + r")-(\d{4})$", acct)
        if mt:
            return f"{mt.group(1)} {mt.group(2)} income tax (est.)"
        if acct in self.labels and not acct.startswith("Assets:Bank:"):
            return self.labels[acct]
        if acct.startswith("Assets:Bank:"):
            parts = acct.split(":")
            kind = strip_mask(parts[-1]).replace("Vault:", "")
            kind = re.sub(r"([a-z])([A-Z])", r"\1 \2", kind).lower()
            return f"{self._inst(acct)} {kind}"
        parts = strip_mask(acct).split(":")
        return " ".join(re.sub(r"([a-z])([A-Z])", r"\1 \2", p) for p in parts[1:])

    def _inst(self, acct):
        for k in sorted(self.labels, key=len, reverse=True):
            if acct.startswith(k + ":") or acct == k:
                return self.labels[k]
        return acct.split(":")[2] if acct.count(":") >= 2 else acct

    def balances(self, on):
        return dict(self.q("select account, sum(amount_minor) from v_ledger where account_type in ('asset','liability') "
                           "and date <= ? group by 1", on))

    def acct_meta(self):
        return {r[0]: r[1:] for r in self.q("select distinct account, report_line, cash_flow_class, account_type from v_ledger")}


# ---------------------------------------------------------------- analysis

def analyse(B, as_of, week_start, baseline, goal, cash_plan, prior=None, open_items=()):
    week_end = as_of - timedelta(days=1)
    A = {"as_of": as_of, "week_start": week_start, "week_end": week_end, "baseline": baseline}
    meta = B.acct_meta()
    b0, b1 = B.balances(baseline.isoformat()), B.balances(as_of.isoformat())
    nw0, nw1 = sum(b0.values()), sum(b1.values())
    A.update(nw0=nw0, nw1=nw1)
    accts = sorted(set(b0) | set(b1))
    A["bs"] = [(a, meta[a][0], b0.get(a, 0), b1.get(a, 0)) for a in accts if b0.get(a, 0) or b1.get(a, 0)]

    # ---- P&L lines in the bridge window (baseline+1 .. as_of), each tagged by what its entry touches
    lines = B.q("""select l.line_id, l.entry_id, l.date, l.account, l.account_type, l.level2, l.amount_minor, l.source,
                          coalesce(l.counterparty, ''), l.description, l.category_source,
                          (select group_concat(distinct o.report_line) from v_ledger o where o.entry_id = l.entry_id
                             and o.statement = 'balance_sheet') as touches
                   from v_ledger l where l.statement = 'income_statement' and l.date > ? and l.date <= ?
                   order by l.date, l.line_id""", baseline.isoformat(), as_of.isoformat())
    bridge = defaultdict(lambda: [0, []])
    for r in lines:
        _, eid, dt, acct, typ, lvl2, amt, src, cp, desc, csrc, touches = r
        touches = touches or ""
        effect = -amt  # P&L line debit-positive: expense +, income -; NW effect is the negative
        if "Crypto" in touches:
            key = "crypto_mtm" if src == "valuation" else "crypto_real"
        elif "Investments" in touches:
            key = "inv_mtm" if src == "valuation" else "inv_real"
        elif typ == "income":
            key = "income"
        elif typ == "expense" and src not in ("accrual", "valuation"):
            key = "spending"
        else:
            key = "other"
        bridge[key][0] += effect
        bridge[key][1].append(r)
    names = {
        "crypto_mtm": "Crypto: market value ([[unrealized|unrealized]])",
        "crypto_real": "Crypto: sold on Kraken ([[realized|realized gain]] less fees)",
        "inv_mtm": "Investments: market value ([[unrealized|unrealized]])",
        "inv_real": "Investments: shares sold ([[realized|realized gain]])",
        "income": "Income (pay, distributions, other)",
        "spending": "Spending",
        "other": "Other (non-cash [[accrual|accruals]], adjustments)",
    }
    pnl_total = sum(v[0] for v in bridge.values())
    gap = (nw1 - nw0) - pnl_total  # balance-sheet-only moves that change net worth (opening/equity entries)
    if gap:
        bridge["other"][0] += gap
    order = ["crypto_mtm", "crypto_real", "inv_mtm", "inv_real", "income", "spending", "other"]
    A["bridge"] = [(k, names[k], bridge[k][0], bridge[k][1]) for k in order if k in bridge or k in ("income", "spending")]
    assert nw0 + sum(x[2] for x in A["bridge"]) == nw1, "bridge does not reconcile"
    A["bridge_gap"] = gap
    A["prior"] = None
    if prior:
        A["prior"] = apply_prior(B, A, prior, b0, b1, bridge, names, meta, baseline, as_of)
    A["open_items"] = list(open_items)

    # ---- week income and spending (Sat..Fri)
    ws, we = week_start.isoformat(), week_end.isoformat()
    spend_sql = """select date, account, level2, amount_minor, coalesce(counterparty, description), description, source
                   from v_ledger where account_type = 'expense' and source not in ('accrual','valuation')
                   and account not like 'Expenses:Fees:Investment%' and date between ? and ? order by date, line_id"""
    wk = B.q(spend_sql, ws, we)
    A["week_spend"] = wk
    A["week_spend_total"] = sum(r[3] for r in wk)
    inc = B.q("""select date, account, -amount_minor, coalesce(counterparty, description), source from v_ledger
                 where account_type = 'income' and date between ? and ? order by date""", ws, we)
    A["week_income"] = inc
    A["week_income_ops"] = sum(r[2] for r in inc if ":Investments:" not in r[1])
    A["week_income_inv"] = sum(r[2] for r in inc if ":Investments:" in r[1])
    t0 = week_start - timedelta(days=28)
    tr = B.q(spend_sql, t0.isoformat(), (week_start - timedelta(days=1)).isoformat())
    A["trail_from"] = t0
    A["trail_total"] = sum(r[3] for r in tr)
    cats = defaultdict(lambda: [0, 0])
    for r in wk:
        cats[r[2]][0] += r[3]
    for r in tr:
        cats[r[2]][1] += r[3]
    A["cats"] = sorted(((k, v[0], v[1] / 4) for k, v in cats.items()), key=lambda x: (-x[1], -x[2]))
    acc_week = B.q("""select account, sum(amount_minor) from v_ledger where account_type='expense' and source='accrual'
                      and date between ? and ? group by 1""", ws, we)
    A["week_accrual_exp"] = acc_week

    # ---- flags
    flags = []
    per = defaultdict(lambda: [None, None, 0])
    for dt, acct, lvl2, amt, cp, desc, src in wk:
        x = per[cp]
        x[0] = x[0] or dt
        x[1] = acct
        x[2] += amt
    for cp, (dt, acct, amt) in per.items():
        prev = B.q("""select date, amount_minor from v_ledger where account_type='expense' and counterparty = ?
                      and date < ? order by date""", cp, ws)
        why = []
        if amt >= CFG["large_charge_dollars"] * 100:
            why.append(f"large (≥ ${CFG['large_charge_dollars']:,})")
        if not prev:
            if amt >= 2000:
                why.append("new payee (first time in the books)")
        else:
            last_d, last_a = prev[-1]
            gap_days = (d(dt) - d(last_d)).days
            recent = [p for p in prev if (d(ws) - d(p[0])).days <= 62]
            first = d(prev[0][0])
            if gap_days >= 300:
                why.append(f"returns after {gap_days // 30} months (last {m(last_a)} on {nice(last_d)}); "
                           "likely an annual renewal or new plan")
            elif (d(ws) - first).days <= 62 and len(recent) >= 1 and len(prev) <= 3:
                why.append(f"new recurring charge (started {nice(prev[0][0])}; {len(prev) + 1} charges so far)")
            if amt >= 2 * last_a and amt >= 10000 and amt - last_a >= 5000 and gap_days < 300:
                why.append(f"up from {m(last_a)} last time")
        if why:
            flags.append((dt, cp, acct, amt, "; ".join(why)))
    flags.sort(key=lambda f: -f[3])
    for k, w, avg in A["cats"]:
        if w - avg >= 10000 and w >= 2 * avg:
            flags.append((we, f"{k} category", "", w, f"week {m(w)} vs average {m(round(avg))}"))
    A["flags"] = flags

    # ---- cash
    cash_accts = [a for a, (rl, cfc, typ) in meta.items() if cfc == "cash" and typ == "asset"]
    cash = [(a, b0.get(a, 0), b1.get(a, 0)) for a in sorted(cash_accts) if b0.get(a, 0) or b1.get(a, 0)]
    A["cash"] = cash
    A["cash_goal_now"] = sum(c[2] for c in cash if c[0] not in B.trips)
    A["cash_goal_then"] = sum(c[1] for c in cash if c[0] not in B.trips)
    if A["prior"]:
        pc = A["prior"]["cash"]
        cash = [(a, pc.get(a, 0), c1) for a, _, c1 in cash]
        A["cash"] = cash
        A["cash_goal_then"] = sum(c[1] for c in cash if c[0] not in B.trips)
    A["trips_deposits"] = B.q("""select date, sum(amount_minor) from v_ledger where account = ? and amount_minor < 0
                                 and date > ? and date <= ? group by 1 order by 1""", TRIPS_RECEIVABLE,
                              baseline.isoformat(), as_of.isoformat())
    A["trips_outflows"] = B.q("""select l.date, coalesce(l.counterparty, l.description), -l.amount_minor from v_ledger l
                                 where l.account in (%s) and l.amount_minor < 0 and l.date > ? and l.date <= ?
                                 order by l.date""" % ",".join("?" * len(B.trips)), *sorted(B.trips),
                              baseline.isoformat(), as_of.isoformat())
    A["crypto_flows"] = B.q("""select coalesce(sum(amount_minor), 0) from v_ledger where account like 'Assets:Crypto:%'
                               and source != 'valuation' and date > ? and date <= ?""", baseline.isoformat(), as_of.isoformat())[0][0]
    A["trips_cash"] = sum(c[2] for c in cash if c[0] in B.trips)
    A["trips_receivable"] = b1.get(TRIPS_RECEIVABLE, 0)
    cards = [(a, -b1.get(a, 0)) for a in sorted(b1) if a.startswith("Liabilities:CreditCard") and b1.get(a, 0)]
    A["cards"] = cards
    A["goal"] = goal
    # cash movements of the week (non-cash legs of entries touching goal cash), largest first
    A["cash_moves"] = B.q("""select coalesce(counterparty, description), sum(cash_effect_minor) from v_cash_flow_lines
                             where date > ? and date <= ? and source != 'pending' group by 1 order by abs(sum(cash_effect_minor)) desc""",
                          baseline.isoformat(), as_of.isoformat())
    A["venmo_pending"] = B.q("""select date, coalesce(counterparty, description), -amount_minor from v_ledger
                                where account = 'Assets:Bank:Venmo' and amount_minor < 0 and date > ? and date <= ?""",
                             baseline.isoformat(), as_of.isoformat())
    A["plan"] = parse_plan(cash_plan)

    # ---- upcoming (as_of .. as_of + 13)
    w_end = as_of + timedelta(days=13)
    A["window"] = (as_of, w_end)
    up = []
    for a, bal in cards:
        pays = B.q("""select date, amount_minor from v_ledger where account = ? and amount_minor > 0
                      and description like '%AUTOPAY%' order by date desc limit 3""", a)
        nxt = ""
        if pays:
            days = [d(p[0]).day for p in pays]
            nxt = f"autopay has run on day {', '.join(str(x) for x in sorted(set(days)))} of the month (last {nice(pays[0][0])}, {m(pays[0][1])})"
        up.append(("card", B.label(a), bal, "Statement balance and due date are not in the books; current balance shown. " + nxt))
    med, later = [], []
    A["later"] = later
    for row in B.q("select accrual_key, description, counterparty, balance_account, open_minor, kind from v_open_payables "
                   "where open_minor != 0"):
        key, desc, cp, acct, amt, kind = row
        if acct.startswith("Liabilities:Accrued:Taxes:Federal") and amt < 0:
            continue  # estimates paid ahead of the accrued amount: nothing owed now
        if "Interest" in acct:
            continue
        if ":Taxes:" in acct or any(x in acct for x in CFG["later_account_substrings"]):
            later.append((B.label(acct), abs(amt)))
            continue  # not due in a 2-week window: taxes, and the payables listed in later_account_substrings
        if ":Medical:" in acct:
            med.append(abs(amt))
            continue
        up.append(("payable", B.label(acct), abs(amt), desc))
    if med:
        up.append(("payable", f"Medical bills ({len(med)}, open)", sum(med), "Member balances from insurance claims; no payment found in connected accounts."))
    for label, amt, note in A["plan"]["outflows"]:
        if re.search(r"this month|Oct", label):
            up.append(("plan", label, amt, f"From the {plan_name(A)} cash plan (your plan, not in the books)."))
    for cp, acct, last_d, last_a, day in recurring(B, as_of, w_end):
        up.append(("recurring", cp, last_a, f"est.: charged monthly around day {day}; last {nice(last_d)}"))
    A["annual"] = annual(B, as_of, w_end)
    for cp, last_d, last_a in A["annual"]:
        up.append(("annual", cp, last_a, f"est.: annual charge; last {nice(last_d)}, so the renewal is due about "
                                         f"{nice(d(last_d).replace(year=d(last_d).year + 1))}"))
    A["upcoming"] = up
    mg = B.q("select date, amount_minor from v_ledger where account like 'Liabilities:Mortgage%' and amount_minor > 0 "
             "and date <= ? order by date desc limit 1", as_of.isoformat())
    A["mortgage"] = mg[0] if mg else None

    # ---- open items
    unc = B.q("select date, account, amount_minor, coalesce(counterparty, description) from v_uncategorized")
    A["uncat"] = unc
    A["unmatched"] = B.q("select date, amount_minor, description from v_ledger where account = 'Assets:Transfers:Unmatched' "
                         "and date <= ? order by date", as_of.isoformat())
    A["pending"] = [(a, -b1[a]) for a in b1 if a.startswith("Liabilities:Pending") and b1[a]]
    A["receivables"] = B.q("select accrual_key, description, counterparty, balance_account, open_minor, basis "
                           "from v_open_receivables where open_minor != 0")
    A["week_lines"] = B.q("""select date, account, amount_minor, coalesce(counterparty,''), description, source
                             from v_ledger where date > ? and date <= ? order by date, entry_id, line_id""",
                          baseline.isoformat(), as_of.isoformat())
    return A


def _has_view(B, name):
    return bool(B.q("select 1 from sqlite_master where name = ?", name))


def plan_name(A):
    mt = re.search(r"cash-plan-(\d{4}-\d{2}-\d{2})", A["plan"].get("path") or "")
    return nice(mt.group(1)).split(" ", 1)[1] if mt else "latest"


def annual(B, as_of, w_end):
    """Yearly charges (at least annual_min_dollars from the report config; software, education, subscriptions,
    entertainment) whose anniversary falls in the
    window and that have not repeated in the last 300 days. Estimates from history; amount = last year's charge."""
    lo, hi = (as_of - timedelta(days=365)).isoformat(), (w_end - timedelta(days=365)).isoformat()
    rows = B.q("""select coalesce(counterparty, description), max(date), sum(amount_minor) from v_ledger
                  where account_type = 'expense' and source = 'connector' and amount_minor >= ? and date between ? and ?
                  and (account like 'Expenses:Software%' or account like 'Expenses:Education%'
                       or account like 'Expenses:Subscriptions%' or account like 'Expenses:Entertainment%')
                  group by 1 order by 2""", CFG["annual_min_dollars"] * 100, lo, hi)
    out = []
    for cp, dt, amt in rows:
        again = B.q("""select 1 from v_ledger where account_type = 'expense' and coalesce(counterparty, description) = ?
                       and amount_minor >= ? and date > ? and date <= ?""", cp, CFG["annual_min_dollars"] * 100,
                    (d(dt) + timedelta(days=60)).isoformat(), as_of.isoformat())
        if not again:
            out.append((cp, dt, amt))
    return out


def apply_prior(B, A, prior, b0, b1, bridge, names, meta, baseline, as_of):
    """Re-anchor the bridge on the prior report's reported balances. Restatement of each account
    (ledger replay at the baseline minus what was reported) is split into: market-value accounts (merged into the
    market lines, since the ledger dates each revaluation on the latest refresh), late-posted P&L entries dated on or
    before the baseline, and any other correction."""
    lab = {}
    for a in set(meta) | set(b0) | set(b1):
        lab.setdefault(B.label(a), a)
    rep_bal, unmapped = defaultdict(int), 0
    for label, v in prior["bs"].items():
        a = lab.get(label)
        if a is None:
            unmapped += v
        else:
            rep_bal[a] += v
    nwR = prior["nw"]
    accts = set(b0) | set(rep_bal)
    restate = {a: b0.get(a, 0) - rep_bal.get(a, 0) for a in accts}
    total = sum(b0.values()) - nwR
    mkt = {"crypto_mtm": sum(v for a, v in restate.items() if a.startswith("Assets:Crypto:")),
           "inv_mtm": sum(v for a, v in restate.items() if a.startswith("Assets:Investments:"))}
    late = B.q("""select l.line_id, l.entry_id, l.date, l.account, l.account_type, l.level2, l.amount_minor, l.source,
                         coalesce(l.counterparty, ''), l.description, l.category_source, ''
                  from v_ledger l join journal_entries e on e.id = l.entry_id
                  where l.statement = 'income_statement' and l.date <= ? and e.created_at > ? and l.source != 'valuation'
                    and not exists (select 1 from v_ledger o where o.entry_id = l.entry_id and
                                    (o.account like 'Assets:Crypto:%' or o.account like 'Assets:Investments:%'))
                  order by l.date, l.line_id""", baseline.isoformat(), prior["cutoff"])
    net = defaultdict(int)
    for r in late:
        net[r[1]] += r[6]
    late = [r for r in late if net[r[1]]]  # drop re-booked entries whose P&L lines net to zero (e.g. payroll splits)
    late_eff = sum(-r[6] for r in late)
    corr = total - mkt["crypto_mtm"] - mkt["inv_mtm"] - late_eff
    for k in ("crypto_mtm", "inv_mtm"):
        if mkt[k]:
            bridge[k][0] += mkt[k]
    names["crypto_mtm"] = "Crypto: market value change since last report ([[unrealized|unrealized]])"
    names["inv_mtm"] = "Investments: market value change since last report ([[unrealized|unrealized]])"
    names["late"] = f"Posted after last report, dated on or before {nice(baseline)} ([[restatement|restatement]])"
    names["restate"] = "Corrections to earlier books ([[restatement|restatement]])"
    if late_eff:
        bridge["late"] = [late_eff, list(late)]
    if corr:
        bridge["restate"] = [corr, []]
    order = ["crypto_mtm", "crypto_real", "inv_mtm", "inv_real", "income", "spending", "late", "restate", "other"]
    A["bridge"] = [(k, names[k], bridge[k][0], bridge[k][1]) for k in order if k in bridge or k in ("income", "spending")]
    ledger_nw0 = A["nw0"]
    A["nw0"] = nwR
    assert nwR + sum(x[2] for x in A["bridge"]) == A["nw1"], "prior-anchored bridge does not reconcile"
    cash = {lab[l]: v for l, v in prior["cash"].items() if l in lab}
    grp = lambda bal, pre: sum(v for a, v in bal.items() if a.startswith(pre))
    return {"nw": nwR, "date": prior["date"], "ledger_nw0": ledger_nw0, "restate_total": total, "mkt": mkt,
            "late": late, "late_eff": late_eff, "corr": corr, "unmapped": unmapped, "cash": cash,
            "crypto0": grp(rep_bal, "Assets:Crypto:"), "crypto1": grp(b1, "Assets:Crypto:"),
            "inv0": grp(rep_bal, "Assets:Investments:"), "inv1": grp(b1, "Assets:Investments:"),
            "dir": prior["dir"]}


def load_prior(path):
    """Reported figures from an earlier weekly.xlsx: net worth (last Bridge row), balance sheet and cash
    (as-of column, by display label), and a cutoff time (the report's file time, UTC) for late-posted entries."""
    from datetime import datetime, timezone
    from openpyxl import load_workbook
    xp = os.path.join(path, "weekly.xlsx")
    wb = load_workbook(xp, data_only=True)
    br = list(wb["Bridge"].iter_rows(values_only=True))
    nw = int(round(br[-1][1] * 100))
    rows = list(wb["Balance Sheet" if "Balance Sheet" in wb.sheetnames else "Balance sheet"].iter_rows(values_only=True))  # sheet renamed Oct 2026
    bs = {r[0]: int(round((r[3] or 0) * 100)) for r in rows[1:] if r[0]}
    crow = list(wb["Cash"].iter_rows(values_only=True))
    cash = {r[0]: int(round((r[2] or 0) * 100)) for r in crow[1:]
            if r[0] and r[3] in ("yes", "no (trips)") and isinstance(r[2], (int, float))}
    assert sum(bs.values()) == nw, f"prior report balance sheet {sum(bs.values())} != reported net worth {nw}"
    src = os.path.join(path, "weekly.html") if os.path.exists(os.path.join(path, "weekly.html")) else xp
    cutoff = datetime.fromtimestamp(os.path.getmtime(src), timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    return {"nw": nw, "bs": bs, "cash": cash, "date": d(str(rows[0][3])), "cutoff": cutoff, "dir": path}


def recurring(B, as_of, w_end):
    """Monthly charges (same payee in each of the last 3 months) whose usual day falls in the window and
    that have not been charged yet this cycle. Estimates from history; amounts = last charge."""
    out = []
    start = (as_of.replace(day=1) - timedelta(days=95)).isoformat()
    rows = B.q("""select coalesce(counterparty, description), account, date, amount_minor from v_ledger
                  where account_type='expense' and source='connector' and date >= ? and date <= ? and amount_minor > 0""",
               start, as_of.isoformat())
    by = defaultdict(list)
    for cp, acct, dt, amt in rows:
        by[cp].append((d(dt), amt, acct))
    for cp, xs in by.items():
        if any(x[2].startswith(("Expenses:Food", "Expenses:Shopping", "Expenses:Transport:Fuel")) for x in xs):
            continue
        months = {(x[0].year, x[0].month) for x in xs}
        prior = [(as_of.replace(day=1) - timedelta(days=30 * k)) for k in (1, 2, 3)]
        if not all((p.year, p.month) in months for p in prior):
            continue
        if len(xs) > 6:  # frequent shopping, not a bill
            continue
        day = round(statistics.median(x[0].day for x in xs))
        try:
            exp = as_of.replace(day=day)
        except ValueError:
            continue
        if exp < as_of:
            continue
        if not (as_of <= exp <= w_end):
            continue
        if any(x[0] >= as_of - timedelta(days=12) for x in xs):
            continue
        last = max(xs)
        out.append((cp, last[2], last[0].isoformat(), last[1], day))
    return sorted(out, key=lambda r: r[4])


def parse_plan(path):
    P = {"path": path, "lines": {}, "outflows": []}
    if not path or not os.path.exists(path):
        return P
    txt = open(path).read()
    for mt in re.finditer(r"^- \*\*(.+?):\*\* (.+)$", txt, re.M):
        P["lines"][mt.group(1)] = mt.group(2)
    sec = re.search(r"## Next 90 days.*?\n\n(.*?)\n\n", txt, re.S)
    if sec:
        for line in sec.group(1).splitlines():
            cells = [c.strip() for c in line.strip("|").split("|")]
            if len(cells) == 2 and cells[1].startswith("$"):
                P["outflows"].append((cells[0], int(round(float(cells[1].strip("$").replace(",", "")) * 100)), ""))
    return P


# ---------------------------------------------------------------- html

def td(c, cls=""):
    neg = "neg" if c < 0 else ""
    return f'<td class="{neg} {cls}">{m(c)}</td>'


def table(head, body):
    return S.table(head, body)


def row(label, vals, cls=""):
    return f'<tr class="{cls}"><td>{label}</td>' + "".join(td(v) if isinstance(v, int) else f"<td>{v}</td>" for v in vals) + "</tr>"


def prior_notes(A, P, bd):
    e = html.escape
    n = []
    if P["crypto0"] or P["crypto1"]:
        flows = A["crypto_flows"]
        n.append(f"Crypto: holdings were {m(P['crypto0'])} in last week's report and {m(P['crypto1'])} today "
                 f"(on-chain value at each refresh's prices), {m(P['crypto1'] - P['crypto0'], True)}. "
                 + ("No coins were bought, sold or moved this week, so the whole change is price." if flows == 0 else
                    f"Coins moved in or out this week ({m(flows, True)} at book value); the rest is price."))
    if P["inv0"] or P["inv1"]:
        n.append(f"{CFG['brokerage_name']} accounts: {m(P['inv0'])} in last week's report, {m(P['inv1'])} today "
                 f"({m(P['inv1'] - P['inv0'], True)}), as reported by {CFG['brokerage_name']} through the bank link.")
    if P["late"]:
        n.append(f"Posted after last week's report but dated on or before {nice(A['baseline'])}: "
                 + ", ".join(f"{e(strip_mask(r[8] or r[9]))} {m(r[6])} ({nice(r[2])})" for r in P["late"])
                 + ". These are real spending; they just arrived late, so they were not in last week's figure.")
    mk = P["mkt"]["crypto_mtm"] + P["mkt"]["inv_mtm"]
    if mk:
        n.append(f"How this is measured (one-time note): the books post each night's market revaluation on the latest "
                 f"day only, so replaying the ledger to {nice(A['baseline'])} now shows crypto and Fidelity closer to cost "
                 f"(net worth {m(P['ledger_nw0'])}, {m(P['restate_total'], True)} above what was reported). This report starts "
                 f"from last week's reported {m(P['nw'])}, so that shift is not shown as a weekly gain or loss.")
    if P["corr"]:
        n.append(f"Corrections to earlier books: {m(P['corr'], True)} of changes to balances dated on or before "
                 f"{nice(A['baseline'])} that are not market moves or late postings (see weekly.xlsx).")
    return n


def build_html(A, B, theme=""):
    T = S.Terms(GLOSS, interactive=not theme)
    e = html.escape
    out = [f'<!doctype html><html lang="en" class="{theme}"><head><meta charset="utf-8">'
           '<meta name="viewport" content="width=device-width, initial-scale=1">'
           f'<title>Weekly report {A["as_of"]}</title><style>{S.CSS}</style></head><body><main>']
    ws, we, ao = A["week_start"], A["week_end"], A["as_of"]
    out.append(f'<h1>Weekly report</h1><p class="sub">{e(CFG["household"])} · week of {nice(ws)} – {nice(we)}, {we.year} · '
               f'books as of {nice(ao)} (nightly refresh)</p>')
    ch = A["nw1"] - A["nw0"]
    gap = A["goal"] * 100 - A["cash_goal_now"]
    goal_k = f"${A['goal'] / 1000:.0f}k"
    NT = CFG["notes"]
    out.append(f'<div class="kpi"><div class="lbl">{T("[[networth|Net worth]]")} now</div><div class="big">{m(A["nw1"])}</div>'
               f'<div class="lbl">{m(ch, True)} vs {"last week&#39;s report, " if A["prior"] else ""}{nice(A["baseline"])} ({m(A["nw0"])})</div></div>')
    out.append('<div class="grid">'
               f'<div class="kpi"><div class="lbl">Spending this week</div><div class="big">{m(A["week_spend_total"])}</div>'
               f'<div class="lbl">avg {m(round(A["trail_total"] / 4))}/wk</div></div>'
               f'<div class="kpi"><div class="lbl">{T("[[cash|Cash]]")} vs {goal_k} goal</div><div class="big">{m(A["cash_goal_now"])}</div>'
               f'<div class="lbl">{m(gap)} to go</div></div></div>')

    # 1. bridge
    out.append(f'<section><h2>1. Net worth {T("[[bridge|bridge]]")}</h2>')
    body = [row(f"Net worth {nice(A['baseline'])}" + (" (last week's report)" if A["prior"] else ""), [A["nw0"]], "total")]
    for k, name, v, _ in A["bridge"]:
        body.append(row(T(name), [v]))
    body.append(row(f"Net worth {nice(ao)}", [A["nw1"]], "grand"))
    out.append(table(["", "Amount"], body))
    notes = []
    bd = {k: (v, ls) for k, _, v, ls in A["bridge"]}
    P = A["prior"]
    if P:
        notes.extend(prior_notes(A, P, bd))
    if "crypto_mtm" in bd and not P and NT["crypto_first_valuation"]:
        notes.append(RC.fill(NT["crypto_first_valuation"], crypto_mtm=m(bd["crypto_mtm"][0])))
    if "inv_real" in bd and not P and NT["inv_real"]:
        notes.append(RC.fill(NT["inv_real"], inv_real=m(bd["inv_real"][0])))
    if "crypto_real" in bd and not P and NT["crypto_real"]:
        notes.append(RC.fill(NT["crypto_real"], crypto_real=m(bd["crypto_real"][0])))
    inc_lines = [r for r in bd.get("income", (0, []))[1]]
    if inc_lines:
        notes.append("Income: " + "; ".join(f"{e(strip_mask(r[8] or r[9]))} {m(-r[6])} ({nice(r[2])})" for r in inc_lines) + ".")
    if A["week_income_ops"] == 0:
        notes.append(NT["no_income"])
    if A["bridge_gap"]:
        notes.append(f"Other includes {m(A['bridge_gap'])} of balance-sheet-only adjustments.")
    notes.append("The bridge adds up exactly to the change in net worth (difference $0.00).")
    out.append('<ul class="notes">' + "".join(f"<li>{T(n, escape=False) if '[[' in n else n}</li>" for n in notes) + "</ul></section>")

    # 2. income and spending
    out.append('<section><h2>2. Profit and Loss, this week</h2>')
    out.append(table(["", "This week"], [
        row("Pay and other income", [A["week_income_ops"]]),
        (row(T("Market value change, crypto and Fidelity ([[realized|realized]] + [[unrealized|unrealized]], same as the bridge)"),
             [sum(v for k, _, v, _ in A["bridge"] if k in ("crypto_mtm", "crypto_real", "inv_mtm", "inv_real"))]) if A["prior"] else
         row(T("Investment gains ([[realized|realized]] + [[unrealized|unrealized]])"), [A["week_income_inv"]])),
        row("Spending (excl. investment fees)", [-A["week_spend_total"]], "total")]))
    out.append(f'<p class="sub" style="margin-top:12px">Spending by category vs {T("[[trailing|4-week average]]")} '
               f'({nice(A["trail_from"])} – {nice(A["week_start"] - timedelta(days=1))})</p>')
    body = [row(e(k), [w, round(a)]) for k, w, a in A["cats"] if w or a >= 500]
    small = [(w, a) for k, w, a in A["cats"] if not (w or a >= 500)]
    if small:
        body.append(row("Other categories (none this week)", [0, round(sum(a for _, a in small))], "sub"))
    body.append(row("Total", [A["week_spend_total"], round(A["trail_total"] / 4)], "total"))
    out.append(table(["Category", "Week", "Avg/wk"], body))
    out.append('<p class="sub" style="margin-top:12px">Flagged items</p><ul class="take">' +
               "".join(f"<li><b>{e(cp)}</b> {m(a)} ({nice(dt)}): {e(why)}</li>" for dt, cp, ac, a, why in A["flags"]) + "</ul>")
    if A["week_accrual_exp"]:
        out.append('<ul class="notes"><li>' + T("Non-cash [[accrual|accrued]] costs dated in the week (not in spending): ") +
                   ", ".join(f"{e(B.label(a))} {m(v)}" for a, v in A["week_accrual_exp"]) + "</li></ul>")
    out.append(f'<ul class="notes"><li>{NT["travel_skew"]}</li></ul></section>' if NT["travel_skew"] and
               any(k == "Travel" and a > 50000 for k, w, a in A["cats"]) else "</section>")

    # 3. cash
    out.append(f'<section><h2>3. Cash vs the ${A["goal"]:,} goal</h2>')
    body = [row(e(B.label(a)) + (" †" if a in B.trips else ""), [c0, c1], "sub" if a in B.trips else "") for a, c0, c1 in A["cash"]]
    body.append(row("Cash counted toward goal", [A["cash_goal_then"], A["cash_goal_now"]], "total"))
    out.append(table(["Account", nice(A["baseline"]), nice(ao)], body))
    cards = sum(b for _, b in A["cards"])
    goal = A["goal"] * 100
    pct = A["cash_goal_now"] / goal * 100
    plan = A["plan"]["lines"]
    out.append('<ul class="take">'
               f"<li>Goal {m(goal)}; counted cash {m(A['cash_goal_now'])} ({pct:.0f}%), <b>{m(gap)} to go</b>. "
               f"Change this week {m(A['cash_goal_now'] - A['cash_goal_then'], True)}.</li>"
               f"<li>{T('[[liquidity|Spendable]]')} after the card balance ({m(cards)}): {m(A['cash_goal_now'] - cards)}.</li>"
               "<li>" + RC.fill(NT["trips_held"], trips_cash=m(A["trips_cash"]), trips_receivable=m(A["trips_receivable"]))
               + (" This week: " + ", ".join(f"{m(-v)} of held paychecks deposited ({nice(dt)})" for dt, v in A["trips_deposits"])
                  + "".join(f"; {e(strip_mask(cp))} {m(v)} paid from the trips account ({nice(dt)})" for dt, cp, v in A["trips_outflows"])
                  + "." if A["trips_deposits"] or A["trips_outflows"] else "") + "</li>")
    if plan:
        pn = plan_name(A)
        out.append(f"<li>Projected date ({pn} cash plan; estimates, not recomputed): "
                   + "; ".join(f"<b>{e(k)}</b> {e(re.sub(r'[*]', '', v))}" for k, v in plan.items()
                               if k.startswith(("Fastest", "No asset", "No changes"))) + ".")
        P = A["prior"]
        if P and any(k.startswith("Fastest") and "KAS" in k for k in plan) and P["crypto1"] and A["crypto_flows"] == 0:
            out.append(f" The fastest path assumes selling the KAS, which has not happened; the wallets are worth "
                       f"{m(P['crypto1'])} today, so that path would now land lower than the plan says.")
        out.append("</li>")
    mv = [x for x in A["cash_moves"] if abs(x[1]) >= 10000][:5]
    out.append("<li>Biggest cash moves: " + ", ".join(f"{e('Amex payment' if cp == 'self' else strip_mask(cp))} {m(v, True)}" for cp, v in mv) + ".</li>")
    low = [(a, c1) for a, _, c1 in A["cash"] if "Checking" in a and 0 < c1 < 100000 and c1 < _ ]
    for a, c1 in low:
        out.append(f"<li>{e(B.label(a))} is down to {m(c1)}; check upcoming autopays drawn from it.</li>")
    out.append("</ul></section>")

    # 4. upcoming
    a0, a1 = A["window"]
    out.append(f'<section><h2>4. Coming up ({nice(a0)} – {nice(a1)})</h2>')
    kind_lbl = {"card": "Card", "payable": "[[ap|Bill owed]]", "plan": "Plan", "recurring": "Usual monthly (est.)",
                "annual": "Yearly renewal (est.)"}
    out.append(table(["Item", "Amount"], [row(T(f"{e(n)} <span class='sub'>· {kind_lbl[k]}</span>", escape=False), [v])
                                          for k, n, v, _ in A["upcoming"]]))
    out.append('<ul class="notes">' + "".join(f"<li>{e(n)}: {e(note)}</li>" for k, n, v, note in A["upcoming"] if note) +
               "<li>Owed but not due in these two weeks: " + "; ".join(f"{e(n)} {m(v)}" for n, v in A["later"]) +
               CFG["later_note"] + "</li>"
               + (f"<li>Mortgage: last payment {nice(A['mortgage'][0])} (principal {m(A['mortgage'][1])}); it posts on the 1st, "
                  "so the next one falls after this window.</li>" if A.get("mortgage") else "") + "</ul></section>")

    # 5. open items
    out.append('<section><h2>5. Needs you</h2><ul class="take">')
    unc = A["uncat"]
    old = all(r[0] < A["week_start"].isoformat() for r in unc)
    out.append(f"<li>{len(unc)} {T('[[uncategorized|uncategorized]]')} lines ({m(sum(abs(r[2]) for r in unc))})"
               + (", all older than this week" if old else "") + "; biggest: " + ", ".join(f"{e(r[3])} {m(abs(r[2]))}" for r in sorted(unc, key=lambda r: -abs(r[2]))[:2]) + ".</li>")
    out.append(f"<li>{len(A['unmatched'])} {T('[[unmatched|unmatched transfers]]')} (net {m(sum(r[1] for r in A['unmatched']))}): "
               + ", ".join(f"{e(r[2])} {nice(r[0])}" for r in A["unmatched"]) + ".</li>")
    for a, v in A["pending"]:
        vp = ", ".join(f"{e(cp)} {m(x)} ({nice(dt)})" for dt, cp, x in A["venmo_pending"])
        out.append(f"<li>{T('[[pending|Pending]]')}: {m(v)} of Venmo payments ({vp}) wait for the bank funding to post; "
                   "nothing to do unless it hasn't posted by next week.</li>")
    rec = defaultdict(int)
    for r in A["receivables"]:
        rec[r[3]] += r[4]
    for a, v in rec.items():
        if a in CFG["receivable_notes"]:
            out.append("<li>" + RC.fill(CFG["receivable_notes"][a], amount=m(v)) + "</li>")
        elif a == TRIPS_RECEIVABLE and NT["trips_deposit"]:
            held = A["trips_receivable"]
            if held:
                out.append("<li>" + RC.fill(NT["trips_deposit"], held=m(held)) + "</li>")
    for cp, last_d, last_a in A.get("annual", []):
        out.append(f"<li>{e(cp)} renews about {nice(d(last_d).replace(year=d(last_d).year + 1))} ({m(last_a)} last year); "
                   "cancel before then if you don't use it.</li>")
    for x in A.get("open_items", []):
        out.append(f"<li>{e(x)}</li>")
    for dt, cp, ac, a_, why in A["flags"]:
        if "returns after" in why or "new recurring" in why:
            out.append(f"<li>{e(cp)} {m(a_)}: {e(why.split(';')[0].split(' (')[0])}; confirm it is expected</li>")
    out.append("</ul></section>")
    if A.get("charts"):
        out.append('<section class="keep"><h2>Trends</h2>' + A["charts"] +
                   '<ul class="notes"><li>Month-end balances and monthly totals from the ledger, checked to the cent '
                   'against the Balance Sheet and Profit and Loss Statement. Full set: scripts/charts.py.</li></ul></section>')
    out.append('<ul class="notes"><li>Source: ledger only (v_ledger, v_open_payables, v_open_receivables, v_uncategorized); '
               f"{CFG['plan_items_phrase']} come from the {plan_name(A)} cash plan and are estimates"
               + (f"; last week's reported figures from {e(os.path.basename(A['prior']['dir']))}/weekly.xlsx" if A["prior"] else "")
               + ". Full detail in weekly.xlsx.</li></ul>")
    if theme:
        out.append('<section class="page"><h2>Glossary</h2>' + T.glossary_html() + "</section>")
    else:
        out.append('<details class="gloss"><summary>Glossary</summary>' + T.glossary_html() + "</details>")
        out.append(f"<script>{S.TIP_JS}</script>")
    out.append("</main></body></html>")
    return "".join(out)


# ---------------------------------------------------------------- xlsx

def build_xlsx(path, A, B):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows, money_cols=()):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(list(r))
        for col in money_cols:
            for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for x in cell:
                    x.number_format = '#,##0.00;(#,##0.00)'
        for i, h in enumerate(header, 1):
            ws.column_dimensions[chr(64 + i)].width = max(14, min(60, len(str(h)) + 4))
        ws.column_dimensions["A"].width = 44
        return ws

    wb.remove(wb.active)
    br = [(f"Net worth {A['baseline']}", A["nw0"] / 100)] + [(S.plain(n), v / 100) for _, n, v, _ in A["bridge"]] + \
         [(f"Net worth {A['as_of']}", A["nw1"] / 100)]
    sheet("Bridge", ["Line", "USD"], br, (2,))
    sheet("Bridge detail", ["Bridge line", "Date", "Account", "Net worth effect", "Counterparty", "Description", "Source"],
          [(S.plain(n), r[2], strip_mask(r[3]), -r[6] / 100, r[8], strip_mask(r[9])[:120], r[7]) for _, n, _, ls in A["bridge"] for r in ls], (4,))
    sheet("Balance Sheet", ["Account", "Line", str(A["baseline"]), str(A["as_of"])],
          [(B.label(a), rl, v0 / 100, v1 / 100) for a, rl, v0, v1 in A["bs"]], (3, 4))
    sheet("Spending by category", ["Category", "This week", "Avg/week (4 wk)"],
          [(k, w / 100, a / 100) for k, w, a in A["cats"]], (2, 3))
    sheet("Week spending", ["Date", "Account", "Amount", "Payee", "Description"],
          [(r[0], strip_mask(r[1]), r[3] / 100, r[4], r[5]) for r in A["week_spend"]], (3,))
    sheet("Flags", ["Date", "Item", "Amount", "Why"], [(r[0], r[1], r[3] / 100, r[4]) for r in A["flags"]], (3,))
    sheet("Cash", ["Account", str(A["baseline"]), str(A["as_of"]), "Counts toward goal"],
          [(B.label(a), c0 / 100, c1 / 100, "no (trips)" if a in B.trips else "yes") for a, c0, c1 in A["cash"]] +
          [("Total toward goal", A["cash_goal_then"] / 100, A["cash_goal_now"] / 100, ""),
           ("Goal", A["goal"], A["goal"], ""), (CFG["trips_receivable_label"], "", A["trips_receivable"] / 100, "no (trips)")], (2, 3))
    sheet("Upcoming", ["Kind", "Item", "Amount", "Note"], [(k, n, v / 100, note) for k, n, v, note in A["upcoming"]], (3,))
    sheet("Open items", ["Type", "Date", "Item", "Amount"],
          [("uncategorized", r[0], r[3], abs(r[2]) / 100) for r in A["uncat"]] +
          [("unmatched transfer", r[0], r[2], r[1] / 100) for r in A["unmatched"]] +
          [("pending", str(A["as_of"]), B.label(a), v / 100) for a, v in A["pending"]] +
          [("receivable", "", r[1][:100], r[4] / 100) for r in A["receivables"]], (4,))
    sheet("All lines", ["Date", "Account", "Amount (debit +)", "Counterparty", "Description", "Source"],
          [(r[0], strip_mask(r[1]), r[2] / 100, r[3], strip_mask(r[4])[:120], r[5]) for r in A["week_lines"]], (3,))
    if A["prior"]:
        P = A["prior"]
        sheet("Restatement", ["Item", "USD"], [
            (f"Net worth {A['baseline']} as reported last week", P["nw"] / 100),
            (f"Net worth {A['baseline']} replayed from today's ledger", P["ledger_nw0"] / 100),
            ("Difference", P["restate_total"] / 100),
            ("  of which crypto (revaluation dated on latest refresh)", P["mkt"]["crypto_mtm"] / 100),
            ("  of which investments (revaluation dated on latest refresh)", P["mkt"]["inv_mtm"] / 100),
            ("  of which late-posted entries", P["late_eff"] / 100),
            ("  of which other corrections", P["corr"] / 100)], (2,))
    sheet("Glossary", ["Term", "Definition"], [tuple(v) for k, v in GLOSS.items()])
    wb.save(path)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--as-of", required=True)
    ap.add_argument("--week-start", required=True)
    ap.add_argument("--baseline", default="", help="baseline date (default: day before week-start)")
    ap.add_argument("--db", default=os.environ.get("FINANCE_DB", os.path.join(DATA, "finance.db")))
    ap.add_argument("--out", default="")
    ap.add_argument("--labels", default=os.path.join(DATA, "config", "labels.yaml"))
    ap.add_argument("--cash-plan", default="", help="cash-plan.md (default: newest exports/cash-plan-*)")
    ap.add_argument("--goal", type=int, default=CFG["goal"])
    ap.add_argument("--trips", default=TRIPS_DEFAULT, help="comma-separated trips-fund accounts")
    ap.add_argument("--chrome", default="")
    ap.add_argument("--prior-report", default="auto", help="earlier weekly-* folder whose reported figures anchor the "
                    "bridge (default: newest earlier one; 'none' to replay the ledger)")
    ap.add_argument("--open-item", action="append", default=[], help="extra 'Needs you' line (from conversation, not the ledger)")
    a = ap.parse_args()
    as_of, ws = d(a.as_of), d(a.week_start)
    global NICE_YEAR
    NICE_YEAR = as_of.year
    base = d(a.baseline) if a.baseline else ws - timedelta(days=1)
    out = a.out or os.path.join(DATA, "exports", f"weekly-{as_of}")
    os.makedirs(out, exist_ok=True)
    plan = a.cash_plan
    if not plan:
        ex = os.path.join(DATA, "exports")
        cands = sorted(x for x in os.listdir(ex) if x.startswith("cash-plan-") and x <= f"cash-plan-{as_of}")
        plan = os.path.join(ex, cands[-1], "cash-plan.md") if cands else ""
    B = Book(a.db, load_labels(a.labels), set(a.trips.split(",")))
    prior = None
    if a.prior_report != "none":
        pr = a.prior_report
        if pr == "auto":
            ex = os.path.join(DATA, "exports")
            c2 = sorted(x for x in os.listdir(ex) if re.fullmatch(r"weekly-\d{4}-\d{2}-\d{2}", x) and x < f"weekly-{as_of}"
                        and os.path.exists(os.path.join(ex, x, "weekly.xlsx")))
            pr = os.path.join(ex, c2[-1]) if c2 else ""
        if pr:
            prior = load_prior(pr)
            if not a.baseline:
                base = prior["date"]
    A = analyse(B, as_of, ws, base, a.goal, plan, prior, a.open_item)
    import charts as C  # trends section (cash vs goal, Profit and Loss by month); validated against the views
    A["charts"] = C.weekly_section(a.db, as_of, a.goal, a.trips.split(",")[0])
    open(os.path.join(out, "weekly.html"), "w").write(build_html(A, B))
    xp = os.path.join(out, "weekly.xlsx")
    build_xlsx(xp, A, B)
    S.reproducible_xlsx(xp, str(as_of))
    chrome = S.find_chrome(a.chrome)
    rc = 0
    if not chrome:
        print("weekly: no Chrome; PDFs skipped", file=sys.stderr)
    else:
        for theme, name in (("light", "weekly.pdf"), ("dark", "weekly-dark.pdf")):
            pdf = os.path.join(out, name)
            fd, tmp = tempfile.mkstemp(prefix=f".weekly-{theme}-", suffix=".html", dir=out)
            with os.fdopen(fd, "w") as f:
                f.write(build_html(A, B, theme=theme))
            try:
                if os.path.exists(pdf):
                    os.remove(pdf)
                r = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                                    f"--print-to-pdf={pdf}", "file://" + os.path.abspath(tmp)],
                                   capture_output=True, text=True, timeout=120)
                if r.returncode or not os.path.exists(pdf):
                    print(f"weekly: {name} failed: {r.stderr[-300:]}", file=sys.stderr)
                    rc = 1
                else:
                    S.reproducible_pdf(pdf, str(as_of))
            finally:
                os.remove(tmp)
    print(f"weekly: net worth {m(A['nw0'])} -> {m(A['nw1'])} ({m(A['nw1'] - A['nw0'], True)}); "
          f"spending {m(A['week_spend_total'])}; cash toward goal {m(A['cash_goal_now'])}; out {out}")
    for k, n, v, _ in A["bridge"]:
        print(f"  {S.plain(n)}: {m(v)}")
    return rc


if __name__ == "__main__":
    sys.exit(main())
